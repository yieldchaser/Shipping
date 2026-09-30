import sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
import docx

doc_p = Path("corpus/01-brokers/Shipbroking_Source_Parsing_Notes.docx")
if doc_p.exists():
    doc = docx.Document(doc_p)
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    print("Length of docx:", len(full_text))
    for p in doc.paragraphs:
        if "poten" in p.text.lower():
            print("MATCH:", p.text)
else:
    print("docx not found")
